"""
OEM re-run comparison — what the verified OEM set does to movable. Excel + scorecard HTML.

Runs M1→M7 IN-MEMORY (no Gold overwrite, anchor run preserved) under three OEM definitions and
diffs the movable:
  1. Anchor (substring) — the validated baseline ($11.888M IS movable).
  2. Verified · web-only — carve out ONLY the 75 web-confirmed OEMs (citations). Strict.
  3. Verified · web + knowledge — also carve the tightened knowledge-pass OEM calls (per-vendor
     best verdict). Fuller carve-out, lower movable.
The point: a more complete OEM carve-out lowers movable below the anchor — quantifying the
"name-list under-counted OEMs" finding so the client decision is evidence-based.

Run:  python reviews/oem_rerun_review.py
"""
from __future__ import annotations
import html as _html
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from engine.io.local import LocalIO
from engine.core.params import Params, load_seed
from engine.core.normalize.cube_adapter import to_canonical, INDIRECT_MAP
from engine.core.normalize.vendor import normalize_vendors
from engine.core.spend.fact import build_fact_spend
from engine.core.spend.pockets import build_pockets
from engine.core.scan.score import score_scan
from engine.core.opportunities import generate as M7

VERSION = "mro-layer0-v1"
IS_L2 = "L2|MRO>Industrial Supplies"
OUT = "reports/OEM_Rerun_Comparison.xlsx"
SCORECARD_HTML = os.environ.get(
    "SCORECARD_OUT",
    "/private/tmp/claude-501/-Users-tanujgupta-Documents-Claude-Projects-Allison-Transmission---MRO-Analysis/"
    "5da58694-a334-491d-8bd7-934ad2d11645/scratchpad/oem_rerun_scorecard.html")

H1 = Font(bold=True, size=16, color="18181B"); H2 = Font(bold=True, size=12, color="18181B")
HEAD = Font(bold=True, color="FFFFFF"); MONO = Font(name="Consolas"); GREY = Font(color="52525C")
HEAD_FILL = PatternFill("solid", fgColor="6A3EBD"); PASS_FILL = PatternFill("solid", fgColor="ECFEF3")
thin = Side(style="thin", color="E4E4E7"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
R = Alignment(horizontal="right"); C = Alignment(horizontal="center")


def _s(ws, cell, v, font=None, fill=None, align=None, border=False):
    c = ws[cell]; c.value = v
    if font: c.font = font
    if fill: c.fill = fill
    if align: c.alignment = align
    if border: c.border = BORDER
    return c


def classification_dicts(io):
    c = io.read("opp", "vendor_classification")
    # Bookend B (full): best per-vendor verdict (web where verified, else tightened knowledge)
    full_fact = {r["vendor_id"]: bool(r["is_oem"]) for _, r in c.iterrows()}
    full_norm = {r["vendor_id"]: {"is_oem": bool(r["is_oem"]), "capability_class": r["capability_class"]}
                 for _, r in c.iterrows()}
    # Bookend A (web-only): trust only web-confirmed; everything else non-OEM (comprehensive -> no substring leak)
    web_fact = {r["vendor_id"]: (bool(r["is_oem"]) if r["source"] == "web" else False) for _, r in c.iterrows()}
    web_norm = {r["vendor_id"]: {"is_oem": (bool(r["is_oem"]) if r["source"] == "web" else False),
                                 "capability_class": r["capability_class"]} for _, r in c.iterrows()}
    return (web_fact, web_norm), (full_fact, full_norm)


def run_chain(canon, params, cfg, oem_brands, engineered, mat_thresh, class_fact, class_norm, rid):
    cim_vendor = normalize_vendors(canon, cfg, material_threshold=mat_thresh, classification=class_norm)
    fact = build_fact_spend(canon, oem_brands=oem_brands, engineered_segments=engineered,
                            run_id=rid, classification=class_fact)
    pockets = build_pockets(fact, run_id=rid)
    mro_fact = fact[fact["l1_code"] == "L1|MRO"]
    mro_pock = pockets[pockets["l1_code"] == "L1|MRO"]
    scan = score_scan(mro_fact, params, run_id=rid)
    out = M7.generate_opportunities(mro_pock, mro_fact, cim_vendor, params, run_id=rid, scan=scan)
    return fact, out


def metrics(label, fact, out):
    opp = out["opportunity"]
    comm = opp[opp["play_route"].isin(["consolidate", "rfp"])]
    is_comm = comm[comm["l2_code"] == IS_L2]
    f = fact
    is_oem_f = f[(f["l2_code"] == IS_L2) & (f["is_oem"])]
    mro_oem_f = f[(f["l1_code"] == "L1|MRO") & (f["is_oem"])]
    return {
        "label": label,
        "is_oem_vendors": int(is_oem_f["vendor_id"].nunique()),
        "is_oem_spend": round(float(is_oem_f["net_spend_usd"].sum()), 0),
        "is_movable": round(float(is_comm["movable_value"].sum()), 0),
        "is_plays": int(len(is_comm)),
        "mro_oem_spend": round(float(mro_oem_f["net_spend_usd"].sum()), 0),
        "mro_movable": round(float(comm["movable_value"].sum()), 0),
        "mro_plays": int(len(comm)),
    }


def compute():
    io = LocalIO(root="data")
    params = Params.load(io, VERSION)
    cfg = load_seed("config/vendor_capability.yaml")
    oem_brands = cfg.get("oem_brands", [])
    engineered = params.list_("engineered_segments")
    mat_thresh = params.num("tier_leverage_spend")
    canon = to_canonical(io.read_bronze("indirect_cube"), INDIRECT_MAP, "indirect")
    (web_f, web_n), (full_f, full_n) = classification_dicts(io)

    rows = []
    f, out = run_chain(canon, params, cfg, oem_brands, engineered, mat_thresh, None, None, "cmp-anchor")
    rows.append(metrics("Anchor (substring)", f, out))
    f, out = run_chain(canon, params, cfg, oem_brands, engineered, mat_thresh, web_f, web_n, "cmp-web")
    rows.append(metrics("Verified · web-confirmed only", f, out))
    f, out = run_chain(canon, params, cfg, oem_brands, engineered, mat_thresh, full_f, full_n, "cmp-full")
    rows.append(metrics("Verified · web + knowledge", f, out))
    return rows


def main():
    rows = compute()
    anchor = rows[0]
    anchor_ok = abs(anchor["is_movable"] - 11888233) <= 50 and anchor["is_plays"] == 19
    wb = Workbook()
    ws = wb.active; ws.title = "OEM re-run comparison"; ws.sheet_view.showGridLines = False
    _s(ws, "A1", "OEM re-run — movable impact of the verified OEM set", H1)
    _s(ws, "A2", "M1→M7 re-run in-memory under three OEM definitions (Gold anchor preserved). A bigger, "
                 "evidence-backed OEM carve-out lowers movable — quantifying the 'name-list under-counted' finding.", GREY)
    heads = ["OEM definition", "IS OEM vendors", "IS OEM spend", "IS movable", "IS plays",
             "MRO OEM spend", "MRO movable"]
    for i, h in enumerate(heads):
        _s(ws, f"{chr(65+i)}4", h, HEAD, HEAD_FILL, Alignment(horizontal="left", wrap_text=True), True)
    r = 5
    for m in rows:
        _s(ws, f"A{r}", m["label"], None, border=True)
        _s(ws, f"B{r}", m["is_oem_vendors"], MONO, align=R, border=True)
        _s(ws, f"C{r}", f"${m['is_oem_spend']:,.0f}", MONO, align=R, border=True)
        _s(ws, f"D{r}", f"${m['is_movable']:,.0f}", Font(bold=True, name="Consolas"), align=R, border=True)
        _s(ws, f"E{r}", m["is_plays"], MONO, align=C, border=True)
        _s(ws, f"F{r}", f"${m['mro_oem_spend']:,.0f}", MONO, align=R, border=True)
        _s(ws, f"G{r}", f"${m['mro_movable']:,.0f}", Font(bold=True, name="Consolas"), align=R, border=True)
        r += 1
    _s(ws, f"A{r+1}", f"Anchor self-check (reproduces $11,888,233 / 19 plays): {'PASS' if anchor_ok else 'FAIL'}",
       Font(bold=True, color="008234" if anchor_ok else "B42318"))
    for col, w in {"A": 30, "B": 15, "C": 16, "D": 16, "E": 9, "F": 16, "G": 16}.items():
        ws.column_dimensions[col].width = w

    os.makedirs("reports", exist_ok=True)
    wb.save(OUT)
    html = render_html(rows, anchor_ok)
    print(f"wrote {OUT}  (anchor self-check {'PASS' if anchor_ok else 'FAIL'})")
    for m in rows:
        print(f"  {m['label']:32s} IS movable ${m['is_movable']:>12,.0f} ({m['is_plays']} plays, {m['is_oem_vendors']} OEM) | "
              f"MRO movable ${m['mro_movable']:>13,.0f}")
    print(f"wrote {html}")


def render_html(rows, anchor_ok):
    def esc(s): return _html.escape(str(s))
    a = rows[0]
    def pct(m):
        return "" if m["label"].startswith("Anchor") else f" ({(m['mro_movable']/a['mro_movable']-1)*100:+.0f}% vs anchor)"
    body = ""
    for i, m in enumerate(rows):
        cls = "anchor" if i == 0 else ""
        body += (f"<tr class='{cls}'><td>{esc(m['label'])}</td><td class='mono'>{m['is_oem_vendors']}</td>"
                 f"<td class='mono'>${m['is_oem_spend']:,.0f}</td><td class='mono b'>${m['is_movable']:,.0f}</td>"
                 f"<td class='mono c'>{m['is_plays']}</td><td class='mono'>${m['mro_oem_spend']:,.0f}</td>"
                 f"<td class='mono b'>${m['mro_movable']:,.0f}<span class='d'>{pct(m)}</span></td></tr>")
    repl = {
        "%%ROWS%%": body,
        "%%ISA%%": f"${a['is_movable']/1e6:.2f}M", "%%ISW%%": f"${rows[1]['is_movable']/1e6:.2f}M",
        "%%ISF%%": f"${rows[2]['is_movable']/1e6:.2f}M",
        "%%ANCHOR%%": "PASS" if anchor_ok else "FAIL",
    }
    page = _TPL
    for k, v in repl.items():
        page = page.replace(k, v)
    os.makedirs(os.path.dirname(SCORECARD_HTML), exist_ok=True)
    with open(SCORECARD_HTML, "w") as f:
        f.write(page)
    return SCORECARD_HTML


_TPL = """<title>OEM Re-run — movable impact · Navanta Lens Engine</title>
<meta name="description" content="What the verified OEM set does to movable: anchor vs web-confirmed vs web+knowledge, re-run through M1-M7. A fuller OEM carve-out lowers movable.">
<style>
  :root{--ink:#18181b;--ink-2:#52525c;--ink-3:#71717a;--border:#e4e4e7;--line:#f4f4f5;--bg:#fafafa;--card:#fff;
    --ok-bg:#ecfef3;--ok-fg:#008234;--warn-bg:#fffbea;--warn-fg:#9e3900;--brand:#6a3ebd;--brand-50:#f7f2ff;
    --mono:ui-monospace,"SF Mono",Menlo,Consolas,monospace;--sans:system-ui,-apple-system,"Segoe UI",Roboto,Helvetica,Arial,sans-serif;}
  *{box-sizing:border-box} body{margin:0;background:var(--bg);color:var(--ink);font-family:var(--sans);line-height:1.55;}
  .wrap{max-width:1000px;margin:0 auto;padding:40px 32px 64px;}
  .mono{font-family:var(--mono);font-variant-numeric:tabular-nums;text-align:right;} .b{font-weight:700;} .c{text-align:center;}
  .eyebrow{font-size:12px;font-weight:600;letter-spacing:.08em;text-transform:uppercase;color:var(--brand);}
  h1{font-size:30px;font-weight:600;letter-spacing:-.01em;margin:6px 0 4px;} .sub{color:var(--ink-2);font-size:15px;max-width:74ch;}
  header{margin-bottom:22px;} .asof{color:var(--ink-3);font-size:13px;margin-top:6px;}
  .kpis{display:grid;grid-template-columns:repeat(3,1fr);gap:14px;margin:18px 0;}
  .kpi{background:var(--card);border:1px solid var(--border);border-radius:12px;padding:16px;} .kpi .k{font-size:12px;color:var(--ink-2);} .kpi .v{font-family:var(--mono);font-size:22px;font-weight:600;margin-top:6px;}
  .kpi.dim .v{color:var(--ink-3);}
  table{border-collapse:collapse;width:100%;font-size:13px;background:var(--card);border:1px solid var(--border);border-radius:12px;overflow:hidden;margin-top:6px;}
  th{background:#fbfbfc;color:var(--ink-2);font-size:11px;text-transform:uppercase;letter-spacing:.03em;text-align:left;padding:10px 12px;border-bottom:1px solid var(--border);white-space:nowrap;}
  td{padding:10px 12px;border-bottom:1px solid var(--line);} td.mono{font-family:var(--mono);font-variant-numeric:tabular-nums;text-align:right;} tr:last-child td{border-bottom:none;}
  tr.anchor{background:var(--brand-50);} tr.anchor td{font-weight:600;} .d{color:var(--warn-fg);font-size:11px;margin-left:6px;}
  .note{display:flex;gap:13px;background:#f0f9ff;border:1px solid #cfe8f5;border-radius:12px;padding:16px 18px;margin-top:18px;} .note p{margin:0;font-size:13px;color:#0a4f6e;} .note b{color:#05405a;}
  footer{margin-top:28px;padding-top:18px;border-top:1px solid var(--border);font-size:13px;color:var(--ink-3);}
</style>
<div class="wrap">
  <header>
    <div class="eyebrow">Navanta Lens · OEM re-run</div>
    <h1>What the verified OEM set does to movable</h1>
    <div class="sub">The engine re-run through M1→M7 under three OEM definitions (Gold anchor untouched). A more complete, evidence-backed OEM carve-out means <b>less</b> spend is consolidatable — so movable drops below the anchor. This is the "name-list under-counted OEMs" finding, quantified, for the client decision.</div>
    <div class="asof">Industrial Supplies + all MRO · indirect cube · anchor self-check %%ANCHOR%%</div>
  </header>
  <div class="kpis">
    <div class="kpi dim"><div class="k">IS movable — anchor (16 OEM)</div><div class="v">%%ISA%%</div></div>
    <div class="kpi"><div class="k">IS movable — web-confirmed (47 OEM)</div><div class="v">%%ISW%%</div></div>
    <div class="kpi"><div class="k">IS movable — web + knowledge</div><div class="v">%%ISF%%</div></div>
  </div>
  <table>
    <thead><tr><th>OEM definition</th><th class="c">IS OEM vendors</th><th class="c">IS OEM spend</th><th class="c">IS movable</th><th class="c">IS plays</th><th class="c">MRO OEM spend</th><th class="c">MRO movable</th></tr></thead>
    <tbody>%%ROWS%%</tbody>
  </table>
  <div class="note"><span>&#8505;</span>
    <p><b>How to choose.</b> <b>Web-confirmed</b> is the citation-backed set — the most defensible "verified" headline; the knowledge-only adds are a review queue. <b>Web + knowledge</b> applies the tightened classifier's best per-vendor verdict (catches small brand OEMs the web-cap skipped, e.g. Fanuc India) but includes unverified calls. Movable sits in that band. Whichever you adopt, it supersedes the substring anchor in M7 only on your sign-off; the Gold anchor run is untouched here. Part-master spec contestability (a later re-run) would move some carve-out back into movable where an OEM's parts are genuinely multi-source.</p>
  </div>
  <footer>In-memory re-run · build_fact_spend(classification=…) supersedes the substring OEM proxy per vendor · anchor reproduces $11,888,233 / 19 plays.</footer>
</div>
"""


if __name__ == "__main__":
    main()
