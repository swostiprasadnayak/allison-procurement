"""
Build the Stage 0 review Excel from the landing reconciliation report.

Convention (per review standard): every figure shown next to its known-good
target with an explicit PASS/FAIL. Sheets:
  - Scorecard          : per-source rows + spend vs documented, with PASS/FAIL
  - Anchor Preview     : Industrial Supplies subset vs the acceptance anchor
  - Schema Divergence  : indirect vs direct cube column mapping (the combine risk)
  - Column Inventory   : every column landed per source

Run:  python reviews/stage0_review.py
Out:  reports/Stage0_Landing_Review.xlsx
"""
from __future__ import annotations
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

REPORT = "reports/stage0_landing_reconciliation.json"
OUT = "reports/Stage0_Landing_Review.xlsx"

# styles
H1 = Font(bold=True, size=16, color="18181B")
H2 = Font(bold=True, size=12, color="18181B")
HEAD = Font(bold=True, color="FFFFFF")
MONO = Font(name="Consolas")
HEAD_FILL = PatternFill("solid", fgColor="6A3EBD")
PASS_FILL = PatternFill("solid", fgColor="ECFEF3")
FAIL_FILL = PatternFill("solid", fgColor="FFF1F2")
PASS_FONT = Font(bold=True, color="008234")
FAIL_FONT = Font(bold=True, color="A7000F")
GREY = Font(color="52525C")
thin = Side(style="thin", color="E4E4E7")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# Schema divergence map (plan §4.2) — the combine risk surface
DIVERGENCE = [
    ("Concept", "Indirect cube column", "Direct cube column"),
    ("Business unit", "Company (AT/OH)", "Cleaned BU / Purchasing Org"),
    ("Vendor (normalized)", "Cleaned Vendor Name", "Cleaned Supplier Name"),
    ("Parent", "Parent Name", "Cleaned Parent Name"),
    ("Taxonomy L1/L2/L3", "Consol 1 / 2 / 3", "Consolidated L1 / L2 / L3"),
    ("Spend (USD)", "Net Spend", "Total Value (USD)"),
    ("Quantity", "Invoice Quantity", "Total Quantity"),
    ("Country / region", "Cleaned Purchasing Country/Region", "Cleaned Supplier Country / Cleaned Plant Region"),
    ("Item id", "Material ID (NOT a cross-vendor part no.)", "Part Number + Part Commodity {Segment/Family/Class/Code}"),
    ("Payment terms", "Payment Terms", "Payment Terms"),
    ("Customer-directed", "Customer Directed (AOH-only)", "Exclusive Customer Directed"),
]


def _set(ws, cell, value, font=None, fill=None, align=None, border=False):
    c = ws[cell]
    c.value = value
    if font: c.font = font
    if fill: c.fill = fill
    if align: c.alignment = align
    if border: c.border = BORDER
    return c


def _verdict(ws, cell, ok):
    c = ws[cell]
    if ok is None:
        c.value = "—"; c.font = GREY
    else:
        c.value = "PASS" if ok else "FAIL"
        c.font = PASS_FONT if ok else FAIL_FONT
        c.fill = PASS_FILL if ok else FAIL_FILL
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER
    return c


def scorecard(wb, rpt):
    ws = wb.active
    ws.title = "Scorecard"
    ws.sheet_view.showGridLines = False
    _set(ws, "A1", "Navanta Lens Engine — Stage 0: Data Landing Review", H1)
    _set(ws, "A2", "Bronze landing of all source files, reconciled to documented counts / totals. PASS = matches reference.", GREY)

    # source table
    r = 4
    cols = ["Source", "Role", "Rows", "Expected rows", "Rows ✓", "Spend (USD)", "Expected spend", "Spend ✓"]
    for i, h in enumerate(cols):
        c = _set(ws, f"{chr(65+i)}{r}", h, HEAD, HEAD_FILL, Alignment(horizontal="center", wrap_text=True), True)
    r += 1
    for s in rpt["sources"]:
        _set(ws, f"A{r}", s["source"], H2, border=True)
        _set(ws, f"B{r}", s.get("role", ""), GREY, align=Alignment(wrap_text=True), border=True)
        _set(ws, f"C{r}", s["rows"], MONO, align=Alignment(horizontal="right"), border=True)
        _set(ws, f"D{r}", s.get("expected_rows") or "—", MONO, align=Alignment(horizontal="right"), border=True)
        _verdict(ws, f"E{r}", s.get("rows_match"))
        _set(ws, f"F{r}", (f"${s['spend_total']:,.2f}" if "spend_total" in s else "—"), MONO, align=Alignment(horizontal="right"), border=True)
        _set(ws, f"G{r}", (f"${s['expected_spend_total']:,.2f}" if "expected_spend_total" in s else "—"), MONO, align=Alignment(horizontal="right"), border=True)
        _verdict(ws, f"H{r}", s.get("spend_match"))
        r += 1

    widths = {"A": 16, "B": 40, "C": 13, "D": 14, "E": 9, "F": 20, "G": 20, "H": 9}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    return ws


def anchor(wb, rpt):
    ws = wb.create_sheet("Anchor Preview")
    ws.sheet_view.showGridLines = False
    _set(ws, "A1", "Industrial Supplies — Acceptance-Anchor Preview", H1)
    _set(ws, "A2", "Subset of the indirect cube (Consol 1='MRO' AND Consol 2='Industrial Supplies'). This is the Stage-2 gate and the basis of the whole $11.9M anchor — confirming it here proves the column mapping.", GREY)
    isc = rpt.get("industrial_supplies_check") or {}
    _set(ws, "A4", isc.get("filter", ""), MONO)
    rows = [
        ("Metric", "Engine", "Reference", "Verdict"),
        ("Rows", isc.get("rows"), isc.get("expected_rows"), isc.get("rows_match")),
        ("Spend (USD)", isc.get("spend"), isc.get("expected_spend"), isc.get("spend_match")),
        ("Distinct vendors", isc.get("vendors"), isc.get("expected_vendors"), isc.get("vendors_match")),
    ]
    r = 6
    for j, (m, e, ref, v) in enumerate(rows):
        if j == 0:
            for i, h in enumerate([m, e, ref, v]):
                _set(ws, f"{chr(65+i)}{r}", h, HEAD, HEAD_FILL, Alignment(horizontal="center"), True)
        else:
            _set(ws, f"A{r}", m, H2, border=True)
            ev = f"${e:,.2f}" if "Spend" in m else f"{e:,}"
            rv = f"${ref:,.2f}" if "Spend" in m else f"{ref:,}"
            _set(ws, f"B{r}", ev, MONO, align=Alignment(horizontal="right"), border=True)
            _set(ws, f"C{r}", rv, MONO, align=Alignment(horizontal="right"), border=True)
            _verdict(ws, f"D{r}", v)
        r += 1
    for col, w in {"A": 18, "B": 20, "C": 20, "D": 9}.items():
        ws.column_dimensions[col].width = w


def divergence(wb):
    ws = wb.create_sheet("Schema Divergence")
    ws.sheet_view.showGridLines = False
    _set(ws, "A1", "Indirect ↔ Direct Cube — Column Divergence", H1)
    _set(ws, "A2", "The two cubes are at different granularities. Normalization must map both into one cim.* shape. Combine is REVERSIBLE — fall back to keeping separate if it breaks consistency (plan §4.4). Needed immediately for fluids + chemicals.", GREY)
    r = 4
    for j, row in enumerate(DIVERGENCE):
        for i, val in enumerate(row):
            font = HEAD if j == 0 else (H2 if i == 0 else MONO)
            fill = HEAD_FILL if j == 0 else None
            _set(ws, f"{chr(65+i)}{r}", val, font, fill, Alignment(wrap_text=True, vertical="top"), True)
        r += 1
    for col, w in {"A": 22, "B": 42, "C": 52}.items():
        ws.column_dimensions[col].width = w


def columns_inventory(wb, rpt):
    ws = wb.create_sheet("Column Inventory")
    ws.sheet_view.showGridLines = False
    _set(ws, "A1", "Columns landed per source (Bronze, raw)", H1)
    r = 3
    for s in rpt["sources"]:
        _set(ws, f"A{r}", f"{s['source']}  ({s['cols']} cols, {s['rows']:,} rows)", H2)
        r += 1
        for i, col in enumerate(s["columns"]):
            _set(ws, f"A{r}", i + 1, GREY, align=Alignment(horizontal="right"))
            _set(ws, f"B{r}", col, MONO)
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 60


def main():
    with open(REPORT) as f:
        rpt = json.load(f)
    wb = Workbook()
    scorecard(wb, rpt)
    anchor(wb, rpt)
    divergence(wb)
    columns_inventory(wb, rpt)
    os.makedirs("reports", exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
